"""
Crea/actualiza una carpeta con el aglomerado de fotos de todos los negocios ya
validados: una foto representativa por negocio, copiada (no movida) desde
fotos/<CASE ID>_<slug>/ y renombrada como "<CASE ID>_Nombre_Del_Negocio.<ext>".

La carpeta fotos/ con sus subcarpetas por negocio se conserva intacta; esto
solo agrega una vista plana adicional para compartir/revisar rapido.

Prioridad de la foto representativa dentro de cada subcarpeta:
  1. La columna "FOTO PRINCIPAL" del Excel maestro, si esta llena -- es la
     eleccion explicita de quien valido el negocio (deberia ser siempre la que
     mas claramente muestre nombre/logos/baterias; ver skill validar-negocio-baterias).
  2. Cualquier archivo que el equipo haya dejado a mano como ejemplo/foto final
     (nombre contiene "ejemplo" o "final") -- se asume que es la mejor toma.
  3. streetview_1.jpg (o el primer streetview_*.jpg que exista) -- ADVERTENCIA:
     esto es solo un ultimo recurso; el heading automatico de la Static API no
     siempre apunta bien, asi que sin "FOTO PRINCIPAL" explicita esto puede
     elegir una foto que no muestre la evidencia con claridad.
  4. maps_foto_1.jpg (o el primer maps_foto_*.jpg que exista).

Uso:
    python scripts/06_aglomerar_fotos.py
"""
import re
import shutil
import sys
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parent.parent
MASTER_PATH = ROOT / "NEGOCIOS_BATERIAS_ARAÑAS.xlsx"
SHEET_NAME = "ar_neg"
FOTOS_DIR = ROOT / "fotos"
AGLOMERADO_DIR = ROOT / "fotos_aglomeradas"

PRIORITY_MARKERS = ("ejemplo", "final")


def sanitize_name(name: str) -> str:
    cleaned = re.sub(r"[^\w\s]", "", str(name), flags=re.UNICODE)
    words = [w.capitalize() for w in cleaned.split()]
    return "_".join(words) or "Negocio"


def pick_representative_photo(folder: Path, foto_principal: str = None):
    files = [f for f in folder.iterdir() if f.is_file() and f.suffix.lower() in (".jpg", ".jpeg", ".png")]
    if not files:
        return None
    if foto_principal:
        for f in files:
            if f.name.strip().lower() == str(foto_principal).strip().lower():
                return f
    for f in files:
        if any(marker in f.stem.lower() for marker in PRIORITY_MARKERS):
            return f
    streetview = sorted(f for f in files if f.stem.lower().startswith("streetview"))
    if streetview:
        return streetview[0]
    maps_photos = sorted(f for f in files if f.stem.lower().startswith("maps_foto"))
    if maps_photos:
        return maps_photos[0]
    return files[0]


def main():
    wb = openpyxl.load_workbook(MASTER_PATH, data_only=True)
    ws = wb[SHEET_NAME]
    headers = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(headers)}
    nombre_col = col.get("NOMBRE NEGOCIO LIMPIO", col["NOMBRE NEGOCIO"])

    AGLOMERADO_DIR.mkdir(exist_ok=True)

    copied, missing = 0, []
    for row in range(2, ws.max_row + 1):
        case_id = ws.cell(row=row, column=col["CASE ID"]).value
        evidencia = ws.cell(row=row, column=col["EVIDENCIA"]).value
        if case_id is None or not evidencia:
            continue  # solo negocios ya validados

        nombre = ws.cell(row=row, column=nombre_col).value or "negocio"
        foto_principal = ws.cell(row=row, column=col["FOTO PRINCIPAL"]).value if "FOTO PRINCIPAL" in col else None
        matches = sorted(FOTOS_DIR.glob(f"{int(case_id):03d}_*"))
        if not matches:
            missing.append(case_id)
            continue

        photo = pick_representative_photo(matches[0], foto_principal)
        if photo is None:
            missing.append(case_id)
            continue

        dest_name = f"{int(case_id)}_{sanitize_name(nombre)}{photo.suffix.lower()}"
        shutil.copy2(photo, AGLOMERADO_DIR / dest_name)
        copied += 1

    print(f"Aglomerado listo en {AGLOMERADO_DIR.name}/: {copied} fotos.")
    if missing:
        print(f"Sin foto representativa para CASE ID: {missing}")


if __name__ == "__main__":
    main()
