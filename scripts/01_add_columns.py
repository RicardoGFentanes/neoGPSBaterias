"""
Agrega al Excel principal:
  - Columna CASE ID (consecutivo 1..n) como primera columna.
  - Columna NOMBRE NEGOCIO LIMPIO (recupera nombre correcto desde el slug de LINK
    cuando NOMBRE NEGOCIO viene con encoding corrupto).
  - Columnas de validacion (vacias / PENDIENTE), al final:
    VENDE BATERIAS, NOMBRE FACHADA DETECTADO, MARCAS LOGOS DETECTADOS,
    BATERIAS EN MOSTRADOR, EVIDENCIA, FOTOS ARCHIVO, FECHA VALIDACION, VALIDADO POR

No corre automaticamente si el archivo esta abierto en Excel (falla al guardar).
Uso:
    python scripts/01_add_columns.py
"""
import sys
from pathlib import Path
from urllib.parse import unquote

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
EXCEL_PATH = ROOT / "NEGOCIOS_BATERIAS_ARAÑAS.xlsx"
SHEET_NAME = "ar_neg"

VALIDATION_COLUMNS = [
    "VENDE BATERIAS",
    "NOMBRE FACHADA DETECTADO",
    "MARCAS LOGOS DETECTADOS",
    "BATERIAS EN MOSTRADOR",
    "EVIDENCIA",
    "FOTOS ARCHIVO",
    "FECHA VALIDACION",
    "VALIDADO POR",
]


def clean_name_from_link(link: str) -> str:
    """Recupera el nombre del negocio desde el slug del link de Maps."""
    try:
        slug = link.split("/place/")[1].split("/data=")[0]
        return unquote(slug).replace("+", " ")
    except IndexError:
        return ""


def main():
    if not EXCEL_PATH.exists():
        print(f"No se encontro {EXCEL_PATH}")
        sys.exit(1)

    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb[SHEET_NAME]

    headers = [c.value for c in ws[1]]
    if "CASE ID" in headers:
        print("La columna CASE ID ya existe, no se vuelve a insertar. Revisa el archivo manualmente si quieres re-generarla.")
        sys.exit(1)

    nombre_col = headers.index("NOMBRE NEGOCIO") + 1
    link_col = headers.index("LINK") + 1

    # 1. Insertar CASE ID como primera columna
    ws.insert_cols(1)
    ws.cell(row=1, column=1, value="CASE ID")
    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=1, value=row - 1)

    # los indices de columna ya capturados (nombre_col, link_col) ahora estan corridos +1
    nombre_col += 1
    link_col += 1

    # 2. Insertar NOMBRE NEGOCIO LIMPIO justo despues de NOMBRE NEGOCIO
    insert_at = nombre_col + 1
    ws.insert_cols(insert_at)
    ws.cell(row=1, column=insert_at, value="NOMBRE NEGOCIO LIMPIO")
    if link_col >= insert_at:
        link_col += 1
    for row in range(2, ws.max_row + 1):
        original = ws.cell(row=row, column=nombre_col).value or ""
        link = ws.cell(row=row, column=link_col).value or ""
        if "�" in str(original):
            cleaned = clean_name_from_link(str(link)) or original
        else:
            cleaned = original
        ws.cell(row=row, column=insert_at, value=cleaned)

    # 3. Columnas de validacion al final
    start_col = ws.max_column + 1
    for i, col_name in enumerate(VALIDATION_COLUMNS):
        ws.cell(row=1, column=start_col + i, value=col_name)
    vende_col = start_col  # VENDE BATERIAS es la primera de estas
    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=vende_col, value="PENDIENTE")

    wb.save(EXCEL_PATH)
    print(f"Listo. {ws.max_row - 1} filas actualizadas en {EXCEL_PATH.name}.")
    print(f"Columnas nuevas: CASE ID, NOMBRE NEGOCIO LIMPIO, {', '.join(VALIDATION_COLUMNS)}")


if __name__ == "__main__":
    main()
