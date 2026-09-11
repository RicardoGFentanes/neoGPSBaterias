"""
Descarga imagenes de Street View Static API para un rango de CASE ID, usando
LAT/LONG del Excel. Guarda DOS tomas por negocio en fotos/<CASEID>_<slug>/:
  - streetview_1.jpg: fov cerrado (zoom), para que se note con claridad el
    nombre en la fachada, logos de marca y baterias en mostrador/aparador.
  - streetview_2.jpg: fov abierto (contexto), para ubicar el negocio en la
    calle si el zoom corta partes relevantes.

El heading se calcula automaticamente (Google apunta la camara hacia el punto
LAT/LONG desde el pano mas cercano), asi que el zoom (fov mas chico) generalmente
alinea bien con la fachada -- si aun asi no se nota la evidencia, revisar en el
navegador otras fechas/angulos en Maps (ver skill validar-negocio-baterias).

Requiere GOOGLE_MAPS_API_KEY en .env (raiz del proyecto).

Uso:
    python scripts/02_download_streetview.py 1 10
    (descarga CASE ID 1 al 10 inclusive; si ya existen las fotos, las reemplaza)
"""
import os
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl
import requests

ROOT = Path(__file__).resolve().parent.parent
EXCEL_PATH = ROOT / "NEGOCIOS_BATERIAS_ARAÑAS.xlsx"
SHEET_NAME = "ar_neg"
FOTOS_DIR = ROOT / "fotos"
ENV_PATH = ROOT / ".env"

STREETVIEW_URL = "https://maps.googleapis.com/maps/api/streetview"
STREETVIEW_METADATA_URL = "https://maps.googleapis.com/maps/api/streetview/metadata"


def load_api_key() -> str:
    if not ENV_PATH.exists():
        print(f"No existe {ENV_PATH}. Copia .env.example a .env y agrega tu API key.")
        sys.exit(1)
    for line in ENV_PATH.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if line.startswith("GOOGLE_MAPS_API_KEY="):
            key = line.split("=", 1)[1].strip()
            if key:
                return key
    print("GOOGLE_MAPS_API_KEY vacio en .env.")
    sys.exit(1)


def slugify(name: str) -> str:
    normalized = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode("ascii")
    slug = re.sub(r"[^a-zA-Z0-9]+", "_", normalized).strip("_").lower()
    return slug or "negocio"


def main():
    if len(sys.argv) != 3:
        print("Uso: python scripts/02_download_streetview.py <case_id_inicio> <case_id_fin>")
        sys.exit(1)
    start_id, end_id = int(sys.argv[1]), int(sys.argv[2])

    api_key = load_api_key()

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb[SHEET_NAME]
    headers = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(headers)}

    if "CASE ID" not in col:
        print("No existe la columna CASE ID. Corre primero scripts/01_add_columns.py")
        sys.exit(1)

    FOTOS_DIR.mkdir(exist_ok=True)

    for row in range(2, ws.max_row + 1):
        case_id = ws.cell(row=row, column=col["CASE ID"]).value
        if case_id is None or not (start_id <= case_id <= end_id):
            continue

        nombre = ws.cell(row=row, column=col.get("NOMBRE NEGOCIO LIMPIO", col["NOMBRE NEGOCIO"])).value or "negocio"
        lat = ws.cell(row=row, column=col["LAT"]).value
        lng = ws.cell(row=row, column=col["LONG"]).value

        if lat is None or lng is None:
            print(f"CASE ID {case_id}: sin LAT/LONG, se omite.")
            continue

        folder = FOTOS_DIR / f"{int(case_id):03d}_{slugify(str(nombre))}"
        folder.mkdir(parents=True, exist_ok=True)

        meta_resp = requests.get(
            STREETVIEW_METADATA_URL,
            params={"location": f"{lat},{lng}", "key": api_key},
            timeout=15,
        )
        meta = meta_resp.json()
        status = meta.get("status")

        if status != "OK":
            print(f"CASE ID {case_id} ({nombre}): sin imagen de Street View (status={status}).")
            (folder / "SIN_STREETVIEW.txt").write_text(f"status={status}\n{meta}", encoding="utf-8")
            continue

        shots = [("streetview_1.jpg", 50), ("streetview_2.jpg", 90)]
        for filename, fov in shots:
            img_resp = requests.get(
                STREETVIEW_URL,
                params={
                    "size": "640x640",
                    "location": f"{lat},{lng}",
                    "fov": fov,
                    "source": "outdoor",
                    "key": api_key,
                },
                timeout=15,
            )
            img_resp.raise_for_status()
            out_path = folder / filename
            out_path.write_bytes(img_resp.content)

        print(f"CASE ID {case_id} ({nombre}): guardado streetview_1.jpg (zoom) y streetview_2.jpg (contexto) en {folder.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
