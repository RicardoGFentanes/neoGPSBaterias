"""
Genera / actualiza VALIDACION_NEGOCIOS_BATERIAS.xlsx: un Excel ligero y enfocado
con los negocios que YA paso el pipeline de validacion (direccion, coordenadas,
veredicto, evidencia, fotos) mas una columna OBSERVACIONES para que el equipo
escriba notas y correcciones.

Es de ida y vuelta con el Excel maestro (NEGOCIOS_BATERIAS_ARAÑAS.xlsx):
  1. Si este Excel de revision ya existe, primero copia lo que se haya escrito
     en OBSERVACIONES hacia el maestro (por CASE ID), para no perder esas notas
     la proxima vez que se regenere.
  2. Regenera este Excel a partir del maestro (ya con OBSERVACIONES sincronizado).

Correr este script despues de cada lote de validacion (o cuando el equipo haya
agregado observaciones que se quieran preservar).

Uso:
    python scripts/04_generar_excel_revision.py
"""
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
MASTER_PATH = ROOT / "NEGOCIOS_BATERIAS_ARAÑAS.xlsx"
REVIEW_PATH = ROOT / "VALIDACION_NEGOCIOS_BATERIAS.xlsx"
FOTOS_DIR = ROOT / "fotos"
SHEET_NAME = "ar_neg"

REVIEW_COLUMNS = [
    "CASE ID",
    "NOMBRE NEGOCIO",
    "DIRECCION",
    "LAT",
    "LONG",
    "LINK",
    "VENDE BATERIAS",
    "NOMBRE FACHADA DETECTADO",
    "MARCAS LOGOS DETECTADOS",
    "BATERIAS EN MOSTRADOR",
    "EVIDENCIA",
    "FOTOS ARCHIVO",
    "OBSERVACIONES",
    "FECHA VALIDACION",
]

COLUMN_WIDTHS = {
    "CASE ID": 8, "NOMBRE NEGOCIO": 28, "DIRECCION": 36, "LAT": 11, "LONG": 11,
    "LINK": 16, "VENDE BATERIAS": 14, "NOMBRE FACHADA DETECTADO": 28,
    "MARCAS LOGOS DETECTADOS": 22, "BATERIAS EN MOSTRADOR": 16,
    "EVIDENCIA": 50, "FOTOS ARCHIVO": 34, "OBSERVACIONES": 40, "FECHA VALIDACION": 14,
}
WRAP_COLUMNS = {"EVIDENCIA", "OBSERVACIONES", "DIRECCION", "FOTOS ARCHIVO"}


def ensure_observaciones_column(ws, headers):
    if "OBSERVACIONES" in headers:
        return headers
    col = ws.max_column + 1
    ws.cell(row=1, column=col, value="OBSERVACIONES")
    headers.append("OBSERVACIONES")
    return headers


def sync_observaciones_from_review(ws_master, col_master):
    """Preserva las OBSERVACIONES ya escritas en el Excel de revision existente."""
    if not REVIEW_PATH.exists():
        return
    wb_review = openpyxl.load_workbook(REVIEW_PATH)
    ws_review = wb_review.active
    review_headers = [c.value for c in ws_review[1]]
    if "CASE ID" not in review_headers or "OBSERVACIONES" not in review_headers:
        return
    rcol = {h: i + 1 for i, h in enumerate(review_headers)}

    row_by_case_id = {}
    for row in range(2, ws_master.max_row + 1):
        cid = ws_master.cell(row=row, column=col_master["CASE ID"]).value
        if cid is not None:
            row_by_case_id[int(cid)] = row

    synced = 0
    for row in range(2, ws_review.max_row + 1):
        cid = ws_review.cell(row=row, column=rcol["CASE ID"]).value
        obs = ws_review.cell(row=row, column=rcol["OBSERVACIONES"]).value
        if cid is None or not obs:
            continue
        master_row = row_by_case_id.get(int(cid))
        if master_row:
            ws_master.cell(row=master_row, column=col_master["OBSERVACIONES"], value=obs)
            synced += 1
    if synced:
        print(f"Se preservaron {synced} observaciones existentes hacia el Excel maestro.")


def find_photo_paths(case_id, fotos_value):
    if not fotos_value:
        return ""
    if not FOTOS_DIR.exists():
        return fotos_value
    matches = sorted(FOTOS_DIR.glob(f"{int(case_id):03d}_*"))
    if not matches:
        return fotos_value
    folder = matches[0].name
    names = [n.strip() for n in str(fotos_value).split(",") if n.strip()]
    return "; ".join(f"fotos/{folder}/{n}" for n in names)


def main():
    wb_master = openpyxl.load_workbook(MASTER_PATH)
    ws_master = wb_master[SHEET_NAME]
    headers = [c.value for c in ws_master[1]]
    headers = ensure_observaciones_column(ws_master, headers)
    col_master = {h: i + 1 for i, h in enumerate(headers)}

    sync_observaciones_from_review(ws_master, col_master)
    wb_master.save(MASTER_PATH)

    # Releer con data_only para tener valores limpios y columnas consistentes
    wb_master = openpyxl.load_workbook(MASTER_PATH, data_only=True)
    ws_master = wb_master[SHEET_NAME]
    headers = [c.value for c in ws_master[1]]
    col_master = {h: i + 1 for i, h in enumerate(headers)}
    nombre_col = col_master.get("NOMBRE NEGOCIO LIMPIO", col_master["NOMBRE NEGOCIO"])
    direccion_col = col_master.get("DIRECCION LIMPIA", col_master["DIRECCION"])

    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "Validacion"

    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    body_font = Font(name="Arial")
    wrap = Alignment(wrap_text=True, vertical="top")
    link_font = Font(name="Arial", color="0563C1", underline="single")

    for i, col_name in enumerate(REVIEW_COLUMNS, start=1):
        cell = ws_out.cell(row=1, column=i, value=col_name)
        cell.font = header_font
        cell.fill = header_fill

    out_row = 2
    for row in range(2, ws_master.max_row + 1):
        evidencia = ws_master.cell(row=row, column=col_master["EVIDENCIA"]).value
        if not evidencia:
            continue  # solo negocios ya pasados por el pipeline
        case_id = ws_master.cell(row=row, column=col_master["CASE ID"]).value
        fotos_raw = ws_master.cell(row=row, column=col_master["FOTOS ARCHIVO"]).value
        values = {
            "CASE ID": case_id,
            "NOMBRE NEGOCIO": ws_master.cell(row=row, column=nombre_col).value,
            "DIRECCION": ws_master.cell(row=row, column=direccion_col).value,
            "LAT": ws_master.cell(row=row, column=col_master["LAT"]).value,
            "LONG": ws_master.cell(row=row, column=col_master["LONG"]).value,
            "LINK": ws_master.cell(row=row, column=col_master["LINK"]).value,
            "VENDE BATERIAS": ws_master.cell(row=row, column=col_master["VENDE BATERIAS"]).value,
            "NOMBRE FACHADA DETECTADO": ws_master.cell(row=row, column=col_master["NOMBRE FACHADA DETECTADO"]).value,
            "MARCAS LOGOS DETECTADOS": ws_master.cell(row=row, column=col_master["MARCAS LOGOS DETECTADOS"]).value,
            "BATERIAS EN MOSTRADOR": ws_master.cell(row=row, column=col_master["BATERIAS EN MOSTRADOR"]).value,
            "EVIDENCIA": evidencia,
            "FOTOS ARCHIVO": find_photo_paths(case_id, fotos_raw),
            "OBSERVACIONES": ws_master.cell(row=row, column=col_master["OBSERVACIONES"]).value,
            "FECHA VALIDACION": ws_master.cell(row=row, column=col_master["FECHA VALIDACION"]).value,
        }
        for i, col_name in enumerate(REVIEW_COLUMNS, start=1):
            cell = ws_out.cell(row=out_row, column=i, value=values[col_name])
            cell.font = body_font
            if col_name in WRAP_COLUMNS:
                cell.alignment = wrap
        link_col_idx = REVIEW_COLUMNS.index("LINK") + 1
        link_cell = ws_out.cell(row=out_row, column=link_col_idx)
        if link_cell.value:
            link_cell.hyperlink = link_cell.value
            link_cell.value = "Ver en Maps"
            link_cell.font = link_font
        out_row += 1

    for i, col_name in enumerate(REVIEW_COLUMNS, start=1):
        ws_out.column_dimensions[ws_out.cell(row=1, column=i).column_letter].width = COLUMN_WIDTHS.get(col_name, 16)

    ws_out.freeze_panes = "A2"
    total = out_row - 2
    last_col_letter = ws_out.cell(row=1, column=len(REVIEW_COLUMNS)).column_letter
    ws_out.auto_filter.ref = f"A1:{last_col_letter}{out_row - 1}"

    wb_out.save(REVIEW_PATH)
    print(f"Excel de validacion generado: {REVIEW_PATH.name} ({total} negocios validados)")


if __name__ == "__main__":
    main()
