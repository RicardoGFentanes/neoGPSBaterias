"""
Reconstruye la direccion (con acentos/enie correctos) a partir de LAT/LONG
usando la Google Geocoding API, ya que la columna DIRECCION original viene
con encoding corrupto en las 270 filas (ver Diccionario de Datos en Obsidian).

Agrega la columna "DIRECCION LIMPIA" justo despues de DIRECCION. Es idempotente:
si una fila ya tiene DIRECCION LIMPIA, no se vuelve a consultar la API (para no
gastar cuota de mas si se corre varias veces).

Requiere Geocoding API habilitada en el proyecto de la API key (.env).

Uso:
    python scripts/05_reverse_geocode.py            # todas las filas pendientes
    python scripts/05_reverse_geocode.py 1 10        # solo CASE ID 1 al 10
"""
import sys
import time
from pathlib import Path

import openpyxl
import requests

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parent.parent
EXCEL_PATH = ROOT / "NEGOCIOS_BATERIAS_ARAÑAS.xlsx"
SHEET_NAME = "ar_neg"
ENV_PATH = ROOT / ".env"
GEOCODE_URL = "https://maps.googleapis.com/maps/api/geocode/json"


def load_api_key() -> str:
    for line in ENV_PATH.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if line.startswith("GOOGLE_MAPS_API_KEY="):
            key = line.split("=", 1)[1].strip()
            if key:
                return key
    print("GOOGLE_MAPS_API_KEY vacio en .env.")
    sys.exit(1)


def main():
    start_id = end_id = None
    if len(sys.argv) == 3:
        start_id, end_id = int(sys.argv[1]), int(sys.argv[2])
    elif len(sys.argv) != 1:
        print("Uso: python scripts/05_reverse_geocode.py [case_id_inicio case_id_fin]")
        sys.exit(1)

    api_key = load_api_key()

    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb[SHEET_NAME]
    headers = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(headers)}

    if "DIRECCION LIMPIA" not in col:
        insert_at = col["DIRECCION"] + 1
        ws.insert_cols(insert_at)
        ws.cell(row=1, column=insert_at, value="DIRECCION LIMPIA")
        headers = [c.value for c in ws[1]]
        col = {h: i + 1 for i, h in enumerate(headers)}

    done, skipped, failed = 0, 0, 0
    for row in range(2, ws.max_row + 1):
        case_id = ws.cell(row=row, column=col["CASE ID"]).value
        if start_id is not None and (case_id is None or not (start_id <= case_id <= end_id)):
            continue

        if ws.cell(row=row, column=col["DIRECCION LIMPIA"]).value:
            skipped += 1
            continue

        lat = ws.cell(row=row, column=col["LAT"]).value
        lng = ws.cell(row=row, column=col["LONG"]).value
        if lat is None or lng is None:
            continue

        resp = requests.get(GEOCODE_URL, params={"latlng": f"{lat},{lng}", "key": api_key, "language": "es"}, timeout=15)
        data = resp.json()
        if data.get("status") == "OK" and data.get("results"):
            address = data["results"][0]["formatted_address"]
            ws.cell(row=row, column=col["DIRECCION LIMPIA"], value=address)
            done += 1
        else:
            print(f"CASE ID {case_id}: geocoding fallo (status={data.get('status')})")
            failed += 1

        time.sleep(0.05)

    wb.save(EXCEL_PATH)
    print(f"Listo. {done} direcciones reconstruidas, {skipped} ya existian, {failed} fallaron.")


if __name__ == "__main__":
    main()
