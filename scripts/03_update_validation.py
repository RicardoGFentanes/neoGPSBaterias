"""
Escribe resultados de validacion (vende baterias, fachada, marcas, etc.) al Excel,
a partir de un JSON con resultados por CASE ID. Pensado para no editar el Excel
celda por celda: se acumulan resultados en un JSON (uno por lote revisado) y este
script los vuelca de una sola vez.

Formato esperado del JSON (lista de objetos):
[
  {
    "case_id": 1,
    "vende_baterias": "SI",              # SI / NO / PENDIENTE
    "nombre_fachada": "Acumuladores Neza",
    "marcas": "LTH, BOSCH",
    "mostrador": "SI",                    # SI / NO
    "baterias_fotos_maps": "2 LTH negras + 1 Full Power roja apiladas en anaquel",
                                            # que baterias (marca/cantidad/tipo) se
                                            # aprecian en las FOTOS DE MAPS especificamente
                                            # (distinto de "mostrador", que es SI/NO general)
    "evidencia": "Mostrador con ~8 baterias visibles, logo LTH en fachada",
    "fotos": "streetview_1.jpg",
    "foto_principal": "streetview_1.jpg",  # opcional: cual de las fotos en "fotos"
                                            # muestra mejor la evidencia (nombre,
                                            # logos, baterias). La usa 06_aglomerar_fotos.py.
    "validado_por": "Claude"
  },
  ...
]

Uso:
    python scripts/03_update_validation.py resultados_lote1.json
"""
import json
import sys
from datetime import date
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
EXCEL_PATH = ROOT / "NEGOCIOS_BATERIAS_ARAÑAS.xlsx"
SHEET_NAME = "ar_neg"

FIELD_TO_COLUMN = {
    "vende_baterias": "VENDE BATERIAS",
    "nombre_fachada": "NOMBRE FACHADA DETECTADO",
    "marcas": "MARCAS LOGOS DETECTADOS",
    "mostrador": "BATERIAS EN MOSTRADOR",
    "baterias_fotos_maps": "BATERIAS VISIBLES EN FOTOS MAPS",
    "evidencia": "EVIDENCIA",
    "fotos": "FOTOS ARCHIVO",
    "foto_principal": "FOTO PRINCIPAL",
    "validado_por": "VALIDADO POR",
}


def ensure_columns_exist(ws, headers):
    """Agrega al final cualquier columna de FIELD_TO_COLUMN que aun no exista."""
    for col_name in FIELD_TO_COLUMN.values():
        if col_name not in headers:
            ws.cell(row=1, column=ws.max_column + 1, value=col_name)
            headers.append(col_name)
    return headers


def main():
    if len(sys.argv) != 2:
        print("Uso: python scripts/03_update_validation.py <resultados.json>")
        sys.exit(1)

    results_path = Path(sys.argv[1])
    results = json.loads(results_path.read_text(encoding="utf-8"))

    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb[SHEET_NAME]
    headers = [c.value for c in ws[1]]
    col_check = {h: i + 1 for i, h in enumerate(headers)}

    if "CASE ID" not in col_check:
        print("No existe la columna CASE ID. Corre primero scripts/01_add_columns.py")
        sys.exit(1)

    headers = ensure_columns_exist(ws, headers)
    col = {h: i + 1 for i, h in enumerate(headers)}

    row_by_case_id = {}
    for row in range(2, ws.max_row + 1):
        cid = ws.cell(row=row, column=col["CASE ID"]).value
        if cid is not None:
            row_by_case_id[int(cid)] = row

    today = date.today().isoformat()
    updated = 0
    for entry in results:
        case_id = int(entry["case_id"])
        row = row_by_case_id.get(case_id)
        if row is None:
            print(f"CASE ID {case_id} no encontrado, se omite.")
            continue
        for field, col_name in FIELD_TO_COLUMN.items():
            if field in entry and col_name in col:
                ws.cell(row=row, column=col[col_name], value=entry[field])
        if "FECHA VALIDACION" in col:
            ws.cell(row=row, column=col["FECHA VALIDACION"], value=today)
        updated += 1

    wb.save(EXCEL_PATH)
    print(f"Actualizados {updated} registros en {EXCEL_PATH.name}.")


if __name__ == "__main__":
    main()
